"""Profit & Loss, Balance Sheet and reconciliation checks, computed from the journal.

Nothing here involves the LLM: statements are sums over journal lines.
"""
from __future__ import annotations

from collections import defaultdict

import pandas as pd

from books.chart_of_accounts import BS_SECTIONS, CASH, PL_SECTIONS, ChartOfAccounts
from books.config import CompanyConfig, to_cents
from books.ledger import JournalLine


def account_balances(journal: list[JournalLine]) -> dict[str, int]:
    """Net debit-positive balance in cents for every account."""
    bal: dict[str, int] = defaultdict(int)
    for line in journal:
        bal[line.account] += line.amount_cents
    return dict(bal)


def _natural(coa: ChartOfAccounts, code: str, cents: int) -> int:
    """Debit-positive cents -> natural presentation (revenue, liabilities, equity shown positive)."""
    return coa.by_code[code].sign * cents


def _money(cents: int) -> float:
    return round(cents / 100, 2)


def build_statements(journal: list[JournalLine], coa: ChartOfAccounts, config: CompanyConfig,
                     txn_ids: list[str], bank_ending_balance: float | None = None,
                     review_count: int = 0) -> dict:
    bal = account_balances(journal)

    # ---- Profit & Loss -------------------------------------------------
    pl_sections = []
    totals: dict[str, int] = {}
    for section in PL_SECTIONS:
        lines = []
        for a in coa.accounts:
            if a.statement == "PL" and a.section == section and bal.get(a.code, 0) != 0:
                lines.append({"code": a.code, "account": a.name, "amount": _money(_natural(coa, a.code, bal[a.code]))})
        total = sum(to_cents(x["amount"]) for x in lines)
        totals[section] = total
        if lines:
            pl_sections.append({"name": section, "lines": lines, "total": _money(total)})
    net_income = (totals["Revenue"] + totals["Other Income"]
                  - totals["Operating Expenses"] - totals["Other Expenses"])

    pnl = {
        "period_start": str(config.period_start),
        "period_end": str(config.period_end),
        "sections": pl_sections,
        "total_revenue": _money(totals["Revenue"]),
        "total_operating_expenses": _money(totals["Operating Expenses"]),
        "operating_income": _money(totals["Revenue"] - totals["Operating Expenses"]),
        "net_income": _money(net_income),
    }

    # ---- Balance Sheet -------------------------------------------------
    bs_sections = []
    section_totals: dict[str, int] = {}
    for section in BS_SECTIONS:
        lines = []
        for a in coa.accounts:
            if a.statement != "BS" or a.section != section:
                continue
            if bal.get(a.code, 0) == 0 and a.code != CASH:
                continue
            cents = bal.get(a.code, 0)
            # Assets are shown debit-positive (contra-assets come out negative);
            # liabilities and equity credit-positive.
            shown = cents if section in ("Current Assets", "Fixed Assets") else -cents
            lines.append({"code": a.code, "account": a.name, "amount": _money(shown)})
        if section == "Equity":
            lines.append({"code": None, "account": "Net Income (current period)", "amount": _money(net_income)})
        total = sum(to_cents(x["amount"]) for x in lines)
        section_totals[section] = total
        bs_sections.append({"name": section, "lines": lines, "total": _money(total)})

    total_assets = section_totals["Current Assets"] + section_totals["Fixed Assets"]
    total_liabilities = section_totals["Liabilities"]
    total_equity = section_totals["Equity"]
    balance_sheet = {
        "as_of": str(config.period_end),
        "sections": bs_sections,
        "total_assets": _money(total_assets),
        "total_liabilities": _money(total_liabilities),
        "total_equity": _money(total_equity),
        "total_liabilities_and_equity": _money(total_liabilities + total_equity),
    }

    # ---- Checks --------------------------------------------------------
    checks = {}
    trial = sum(bal.values())
    checks["trial_balance_zero"] = {"passed": trial == 0, "detail": f"Sum of all debits and credits = {_money(trial)}"}
    diff = total_assets - (total_liabilities + total_equity)
    checks["balance_sheet_balances"] = {
        "passed": diff == 0,
        "detail": f"Assets {_money(total_assets):,.2f} vs Liabilities + Equity "
                  f"{_money(total_liabilities + total_equity):,.2f} (difference {_money(diff):,.2f})",
    }
    cash = bal.get(CASH, 0)
    if bank_ending_balance is not None:
        cash_diff = cash - to_cents(bank_ending_balance)
        checks["cash_ties_to_bank"] = {
            "passed": cash_diff == 0,
            "detail": f"Balance sheet cash {_money(cash):,.2f} vs bank ending balance "
                      f"{bank_ending_balance:,.2f} (difference {_money(cash_diff):,.2f})",
        }
    posted = {line.txn_id for line in journal if line.txn_id and line.account == CASH}
    expected = set(txn_ids)
    missing, extra = expected - posted, posted - expected
    checks["all_transactions_posted"] = {
        "passed": not missing and not extra,
        "detail": f"{len(posted & expected)} of {len(expected)} posted; missing {sorted(missing)[:5]}, "
                  f"unexpected {sorted(extra)[:5]}",
    }
    checks["no_items_awaiting_review"] = {
        "passed": review_count == 0,
        "detail": f"{review_count} transaction(s) flagged for human review",
    }

    return {
        "business_name": config.business_name,
        "profit_and_loss": pnl,
        "balance_sheet": balance_sheet,
        "checks": checks,
    }


def flatten_statements(result: dict) -> dict:
    """Compact {account name: amount} view used by evaluators and dataset references."""
    pnl = {line["account"]: line["amount"]
           for s in result["profit_and_loss"]["sections"] for line in s["lines"]}
    bs = {line["account"]: line["amount"]
          for s in result["balance_sheet"]["sections"] for line in s["lines"]}
    return {
        "pnl": pnl,
        "bs": bs,
        "net_income": result["profit_and_loss"]["net_income"],
        "total_revenue": result["profit_and_loss"]["total_revenue"],
        "total_assets": result["balance_sheet"]["total_assets"],
        "total_liabilities": result["balance_sheet"]["total_liabilities"],
        "total_equity": result["balance_sheet"]["total_equity"],
        "checks": {k: v["passed"] for k, v in result["checks"].items()},
    }


# ---- Excel export -------------------------------------------------------

def export_excel(path, result: dict, categorized: pd.DataFrame, journal_df: pd.DataFrame,
                 report_dict: dict, loan_schedule: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    money_fmt = '#,##0.00;(#,##0.00);"-"'
    bold = Font(bold=True)
    title_font = Font(bold=True, size=14)
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    header_font = Font(bold=True, color="FFFFFF")
    top_border = Border(top=Side(style="thin"))
    double_border = Border(top=Side(style="thin"), bottom=Side(style="double"))
    review_fill = PatternFill("solid", fgColor="FFF2CC")
    fail_fill = PatternFill("solid", fgColor="F8CBAD")
    pass_fill = PatternFill("solid", fgColor="C6EFCE")

    wb = Workbook()

    def statement_sheet(ws, title, subtitle, sections, closing_rows):
        ws["A1"], ws["A1"].font = result["business_name"], title_font
        ws["A2"], ws["A2"].font = title, bold
        ws["A3"] = subtitle
        r = 5
        for s in sections:
            ws.cell(r, 1, s["name"]).font = bold
            r += 1
            for line in s["lines"]:
                ws.cell(r, 1, f"    {line['account']}")
                ws.cell(r, 2, line["amount"]).number_format = money_fmt
                r += 1
            ws.cell(r, 1, f"Total {s['name']}").font = bold
            c = ws.cell(r, 2, s["total"])
            c.number_format, c.font, c.border = money_fmt, bold, top_border
            r += 2
        for label, value in closing_rows:
            ws.cell(r, 1, label).font = bold
            c = ws.cell(r, 2, value)
            c.number_format, c.font, c.border = money_fmt, bold, double_border
            r += 1
        ws.column_dimensions["A"].width = 44
        ws.column_dimensions["B"].width = 18

    pnl = result["profit_and_loss"]
    ws = wb.active
    ws.title = "Profit & Loss"
    statement_sheet(ws, "Profit & Loss Statement", f"{pnl['period_start']} to {pnl['period_end']}",
                    pnl["sections"], [("Net Income", pnl["net_income"])])

    bs = result["balance_sheet"]
    ws = wb.create_sheet("Balance Sheet")
    statement_sheet(ws, "Balance Sheet", f"As of {bs['as_of']}", bs["sections"], [
        ("Total Assets", bs["total_assets"]),
        ("Total Liabilities + Equity", bs["total_liabilities_and_equity"]),
    ])

    ws = wb.create_sheet("Checks")
    ws.append(["Check", "Result", "Detail"])
    for name, chk in result["checks"].items():
        ws.append([name, "PASS" if chk["passed"] else "FAIL", chk["detail"]])
        ws.cell(ws.max_row, 2).fill = pass_fill if chk["passed"] else fail_fill
    ws.append([])
    ws.append(["Input validation", "OK" if report_dict["ok"] else "FAILED"])
    for e in report_dict["errors"]:
        ws.append(["  error", e])
    for w in report_dict["warnings"]:
        ws.append(["  warning", w])
    for k, v in report_dict["stats"].items():
        ws.append([f"  {k}", v])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 100

    def table_sheet(name, frame, money_cols=(), highlight_col=None):
        ws = wb.create_sheet(name)
        ws.append(list(frame.columns))
        for row in frame.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
        for i, col in enumerate(frame.columns, 1):
            ws.cell(1, i).fill, ws.cell(1, i).font = header_fill, header_font
            width = min(60, max(10, int(frame[col].astype(str).str.len().quantile(0.9)) + 2 if len(frame) else 10,
                                len(str(col)) + 2))
            ws.column_dimensions[get_column_letter(i)].width = width
            if col in money_cols:
                for r in range(2, ws.max_row + 1):
                    ws.cell(r, i).number_format = money_fmt
        if highlight_col and highlight_col in frame.columns:
            ci = list(frame.columns).index(highlight_col) + 1
            for r in range(2, ws.max_row + 1):
                if ws.cell(r, ci).value is True:
                    for c in range(1, len(frame.columns) + 1):
                        ws.cell(r, c).fill = review_fill
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top")

    txn_cols = ["txn_id", "date", "description", "amount", "balance", "category", "confidence",
                "source", "needs_review", "reasoning"]
    txns = categorized[txn_cols].copy()
    table_sheet("Transactions", txns, money_cols=("amount", "balance"), highlight_col="needs_review")

    review = txns[txns["needs_review"]].copy()
    review.insert(6, "corrected_category", "")
    table_sheet("Review Queue", review.drop(columns=["needs_review"]), money_cols=("amount", "balance"))

    table_sheet("Journal", journal_df, money_cols=("debit", "credit"))
    if loan_schedule:
        table_sheet("Loan Schedule", pd.DataFrame(loan_schedule),
                    money_cols=("payment", "interest", "principal", "balance_after"))

    wb.save(path)
