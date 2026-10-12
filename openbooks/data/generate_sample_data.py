"""Generate synthetic 2026 bank data with known-correct categories.

    python data/generate_sample_data.py

Writes (all deterministic, fixed seeds):
    data/sample_transactions_2026.csv    bank-style export (Date, Description, Withdrawals, Deposits, Balance)
    data/sample_transactions_2026.xlsx   same data, with a bank preamble above the header row
    data/sample_truth_2026.csv           txn_id -> correct category (the "answer key")
    data/golden_transactions.csv         labeled categorization eval set (sampled + hand-written hard cases)
    data/scenarios/*.json                full-year scenarios for the statement eval

Replace these with real, accountant-labeled data before trusting eval scores.
"""
from __future__ import annotations

import calendar
import csv
import json
import random
import string
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

DATA_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(DATA_DIR.parent))

YEAR = 2026
CLIENTS = ["ACME CORP", "BLUEWAVE HEALTH", "KESTREL LOGISTICS", "ORBIT FOODS", "SUMMIT LEARNING"]
CITIES = ["BOSTON", "CHICAGO", "NEW YORK", "AUSTIN", "DENVER"]

MAIN_CONFIG = json.loads((DATA_DIR.parent / "config" / "company.json").read_text())


def generate_year(seed: int, opening_cash: float, scale: float = 1.0) -> list[dict]:
    rng = random.Random(seed)
    txns: list[tuple[date, int, str, float, str]] = []

    def money(lo, hi):
        return round(rng.uniform(lo, hi), 2)

    def ref(n=4):
        return "".join(rng.choice(string.ascii_uppercase + string.digits) for _ in range(n))

    gusto = ref(4)
    for m in range(1, 13):
        last = calendar.monthrange(YEAR, m)[1]

        def add(day, desc, amount, category):
            txns.append((date(YEAR, m, min(day, last)), len(txns), desc, round(amount, 2), category))

        def rday(lo=2, hi=28):
            return rng.randint(lo, hi)

        # Revenue
        for _ in range(rng.randint(3, 5)):
            c = rng.choice(CLIENTS)
            if rng.random() < 0.75:
                add(rday(), f"ACH CREDIT {c} INV-{rng.randint(1000, 1999)}", money(5000, 12000) * scale, "Service Revenue")
            else:
                add(rday(), f"WIRE IN {c} REF {rng.randint(10000, 99999)}", money(6000, 14000) * scale, "Service Revenue")
        for _ in range(rng.randint(2, 4)):
            add(rday(), f"STRIPE TRANSFER ST-{ref(6)}", money(800, 3000) * scale, "Service Revenue")
        add(last, "INTEREST PAYMENT", money(4, 9), "Interest Income")

        # Payroll
        for payday in (15, last):
            add(payday, f"GUSTO NET PAY {gusto}", -money(7800, 8200) * scale, "Wages")
            add(payday, f"GUSTO TAX {gusto}", -money(2100, 2300) * scale, "Payroll Taxes")
        add(15, f"GUSTO FEE {gusto}", -80.00, "Professional Fees")

        # Fixed monthly costs
        add(1, "ONLINE PMT RIVERSIDE PROPERTIES LLC", -3200.00, "Rent")
        add(5, "SBA LOAN PMT 4471", -1050.00, "Loan Payable")
        add(3, "ADOBE *CREATIVE CLD ADOBE.LY/ENUS", -179.97, "Software & Subscriptions")
        add(5, "GOOGLE *WORKSPACE NORTHWI", -84.00, "Software & Subscriptions")
        add(7, "SLACK T02KX9 DUBLIN", -43.75, "Software & Subscriptions")
        add(9, "FIGMA MONTHLY", -45.00, "Software & Subscriptions")
        add(11, "ZOOM.US 888-799-9666", -15.99, "Software & Subscriptions")
        add(12, "INTUIT *QBOOKS ONLINE", -90.00, "Software & Subscriptions")
        add(14, "APPLE.COM/BILL 866-712-7753", -9.99, "Software & Subscriptions")
        add(2, "AMAZON WEB SERVICES AWS.AMAZON.CO WA", -money(300, 650), "Cloud Hosting")
        add(4, f"GOOGLE *CLOUD {ref(5)}", -money(40, 120), "Cloud Hosting")
        add(18, "COMCAST BUSINESS 800-391-3000", -189.95, "Utilities")
        add(20, "PEPCO BILL PAY", -money(150, 320), "Utilities")
        add(last, "MONTHLY MAINTENANCE FEE", -25.00, "Bank & Merchant Fees")

        # Variable costs
        for _ in range(rng.randint(1, 3)):
            add(rday(), f"AMAZON MKTPL*{ref(5)} AMZN.COM/BILL WA", -money(25, 180), "Office Supplies")
        if rng.random() < 0.5:
            add(rday(), f"STAPLES 00{rng.randint(1000, 9999)}", -money(40, 120), "Office Supplies")
        if rng.random() < 0.35:
            add(rday(), f"AMAZON MKTPL REFUND*{ref(5)}", money(15, 60), "Office Supplies")
        add(6, f"GOOGLE *ADS{rng.randint(1000, 9999)}", -money(300, 900), "Advertising & Marketing")
        if rng.random() < 0.4:
            add(rday(), f"FACEBK *ADS {ref(6)}", -money(200, 500), "Advertising & Marketing")
        if m in (3, 9):
            add(rday(), "LINKEDIN *JOBS", -295.00, "Advertising & Marketing")
        if rng.random() < 0.4:
            d0 = rday(2, 20)
            city = rng.choice(CITIES)
            add(d0, f"DELTA AIR {rng.randint(10**9, 10**10 - 1)}", -money(350, 700), "Travel")
            add(d0 + 3, f"MARRIOTT {city}", -money(380, 620), "Travel")
            add(d0 + 1, f"SWEETGREEN {city}", -money(14, 32), "Meals")
        for _ in range(rng.randint(1, 4)):
            add(rday(), "UBER *TRIP HELP.UBER.COM", -money(14, 60), "Travel")
        for _ in range(rng.randint(3, 7)):
            desc = rng.choice([f"STARBUCKS STORE {rng.randint(10000, 99999)}", "UBER *EATS 8005928996",
                               f"DOORDASH*{rng.choice(['CHIPOTLE', 'PANERA', 'SHAKE SHACK'])}",
                               f"PANERA BREAD #{rng.randint(100, 999)}"])
            add(rday(), desc, -money(8, 85), "Meals")
        add(22, f"ZELLE TO MARIA LOPEZ MEMO: DESIGN INV {m:02d}{rng.randint(10, 28)}", -money(1500, 3000) * scale,
            "Contractors")
        if rng.random() < 0.4:
            add(rday(), f"UPWORK -REF{rng.randint(100000, 999999)}", -money(400, 1200), "Contractors")
        if rng.random() < 0.25:
            add(rday(), "WIRE TRANSFER FEE", -30.00, "Bank & Merchant Fees")

        # Quarterly
        if m in (1, 4, 7, 10):
            add(10, "THE HARTFORD INS PREM", -1140.00, "Insurance")
            add(25, "SMITH & CO CPA", -950.00, "Professional Fees")
            add(15, f"IRS USATAXPYMT {rng.randint(10**11, 10**12 - 1)}", -3500.00 * scale, "Owner Draws")

        # Balance-sheet movements
        add(28, "ONLINE TRANSFER TO CHK XXXXXX8812 PERSONAL", -4000.00 * scale, "Owner Draws")
        if m in (2, 5, 8, 11):
            add(26, "ONLINE TRANSFER TO SAV XXXXXX3301", -2000.00, "Savings Account")
        if m == 10:
            add(3, "ONLINE TRANSFER FROM SAV XXXXXX3301", 5000.00, "Savings Account")
        if m == 3:
            add(17, "DEPOSIT FROM M PATEL PERSONAL CHK", 10000.00, "Owner Contributions")
        if m == 5:
            add(19, "APPLE STORE #R102", -2899.00, "Equipment")
        if m == 9:
            add(23, "DELL MARKETING L.P.", -1640.00, "Equipment")
        if m == 6:
            add(rday(), "LEGALZOOM*ORDER 888-773-0888", -249.00, "Professional Fees")

    txns.sort(key=lambda t: (t[0], t[1]))
    out, bal = [], round(opening_cash * 100)
    for i, (d, _, desc, amt, cat) in enumerate(txns, 1):
        bal += round(amt * 100)
        out.append({"txn_id": f"T{i:04d}", "date": d.isoformat(), "description": desc,
                    "amount": round(amt, 2), "balance": bal / 100, "category": cat})
    return out


# Hand-written hard cases for the categorization golden set.
HARD_CASES = [
    ("AMAZON WEB SERVICES AWS.AMAZON.CO WA", -412.18, "Cloud Hosting", "AWS vs Amazon Marketplace"),
    ("AMAZON MKTPL*7HG2Q AMZN.COM/BILL WA", -58.40, "Office Supplies", "Amazon Marketplace vs AWS"),
    ("AMAZON MKTPL REFUND*K2J8D", 42.99, "Office Supplies", "Refund keeps the original category"),
    ("UBER *EATS 8005928996", -38.42, "Meals", "Uber Eats is food, not travel"),
    ("UBER *TRIP HELP.UBER.COM", -24.10, "Travel", "Uber trip is travel, not meals"),
    ("GOOGLE *ADS4471", -612.00, "Advertising & Marketing", "Google Ads vs Workspace vs Cloud"),
    ("GOOGLE *WORKSPACE NORTHWI", -84.00, "Software & Subscriptions", "Google Workspace vs Ads"),
    ("GOOGLE *CLOUD 7HX2K", -88.31, "Cloud Hosting", "Google Cloud vs Workspace"),
    ("IRS USATAXPYMT 270612345678", -3500.00, "Owner Draws", "Owner's estimated tax for pass-through LLC"),
    ("ONLINE TRANSFER TO SAV XXXXXX3301", -2000.00, "Savings Account", "Own savings, not an owner draw"),
    ("ONLINE TRANSFER TO CHK XXXXXX8812 PERSONAL", -4000.00, "Owner Draws", "Personal account = draw"),
    ("ONLINE TRANSFER FROM SAV XXXXXX3301", 5000.00, "Savings Account", "Inflow that is not revenue"),
    ("DEPOSIT FROM M PATEL PERSONAL CHK", 10000.00, "Owner Contributions", "Inflow that is not revenue"),
    ("SBA LOAN PMT 4471", -1050.00, "Loan Payable", "Loan payment is not an expense"),
    ("APPLE STORE #R102", -2899.00, "Equipment", "Capitalized equipment, not supplies"),
    ("DELL MARKETING L.P.", -1640.00, "Equipment", "Capitalized equipment"),
    ("APPLE.COM/BILL 866-712-7753", -9.99, "Software & Subscriptions", "Apple subscription vs Apple Store"),
    ("STRIPE TRANSFER ST-K2J8X1", 2315.40, "Service Revenue", "Processor payout is revenue"),
    ("WIRE IN ACME CORP REF 88812", 12500.00, "Service Revenue", "Client wire"),
    ("INTEREST PAYMENT", 6.12, "Interest Income", "Bank interest is other income, not revenue"),
    ("GUSTO FEE 9QX1", -80.00, "Professional Fees", "Payroll processing fee, not wages"),
    ("GUSTO TAX 9QX1", -2214.55, "Payroll Taxes", "Payroll tax vs wages"),
    ("ZELLE TO MARIA LOPEZ MEMO: DESIGN INV 0314", -2400.00, "Contractors", "Payment to individual with invoice memo"),
    ("CHECK #1043", -1200.00, "Uncategorized", "No payee information: should go to review"),
    ("VENMO PAYMENT 1029384756", -150.00, "Uncategorized", "Unexplained P2P payment: should go to review"),
    ("LINKEDIN *JOBS", -295.00, "Advertising & Marketing", "Job post is advertising"),
    ("THE HARTFORD INS PREM", -1140.00, "Insurance", ""),
    ("SMITH & CO CPA", -950.00, "Professional Fees", ""),
    ("WIRE TRANSFER FEE", -30.00, "Bank & Merchant Fees", "Fee, not a transfer"),
    ("DOORDASH*SHAKE SHACK", -27.84, "Meals", ""),
    ("MARRIOTT CHICAGO", -512.70, "Travel", ""),
    ("FACEBK *ADS QW9E2R", -340.00, "Advertising & Marketing", "Meta ads billing descriptor"),
]


def write_bank_csv(rows: list[dict], path: Path) -> None:
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Date", "Description", "Withdrawals", "Deposits", "Balance"])
        for r in rows:
            amt = r["amount"]
            w.writerow([date.fromisoformat(r["date"]).strftime("%m/%d/%Y"), r["description"],
                        f"{-amt:.2f}" if amt < 0 else "", f"{amt:.2f}" if amt > 0 else "", f"{r['balance']:.2f}"])


def write_bank_xlsx(rows: list[dict], path: Path, business: str) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append([f"{business} - Business Checking ...7720"])
    ws.append([f"Statement period 01/01/{YEAR} - 12/31/{YEAR}"])
    ws.append([])
    ws.append(["Posted Date", "Description", "Amount", "Running Balance"])
    for r in rows:
        ws.append([date.fromisoformat(r["date"]), r["description"], r["amount"], r["balance"]])
        ws.cell(ws.max_row, 1).number_format = "mm/dd/yyyy"
        ws.cell(ws.max_row, 3).number_format = "#,##0.00"
        ws.cell(ws.max_row, 4).number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 13
    ws.column_dimensions["D"].width = 16
    wb.save(path)


def build_golden(seed: int = 7, per_category: int = 6) -> list[dict]:
    rng = random.Random(seed)
    year = generate_year(seed, opening_cash=10000)
    by_cat = defaultdict(list)
    for r in year:
        by_cat[r["category"]].append(r)
    hard_descs = {h[0] for h in HARD_CASES}
    golden = []
    for cat in sorted(by_cat):
        pool = [r for r in by_cat[cat] if r["description"] not in hard_descs]
        # Prefer distinct descriptions so the set isn't 4 copies of "SBA LOAN PMT"
        seen, distinct = set(), []
        for r in rng.sample(pool, len(pool)):
            if r["description"] not in seen:
                seen.add(r["description"])
                distinct.append(r)
        for r in distinct[:per_category]:
            golden.append({"date": r["date"], "description": r["description"], "amount": r["amount"],
                           "category": cat, "difficulty": "typical", "note": ""})
    for i, (desc, amt, cat, note) in enumerate(HARD_CASES):
        golden.append({"date": f"{YEAR}-{(i % 12) + 1:02d}-15", "description": desc, "amount": amt,
                       "category": cat, "difficulty": "hard", "note": note})
    for i, g in enumerate(golden, 1):
        g["txn_id"] = f"G{i:03d}"
    return golden


def scenario(name: str, description: str, seed: int, config: dict, scale: float) -> dict:
    rows = generate_year(seed, opening_cash=config["opening_balances"]["1000"], scale=scale)
    return {
        "name": name,
        "description": description,
        "config": config,
        "transactions": [{k: r[k] for k in ("txn_id", "date", "description", "amount", "balance")} for r in rows],
        "truth": {r["txn_id"]: r["category"] for r in rows},
    }


def main() -> None:
    cfg = MAIN_CONFIG
    rows = generate_year(seed=2026, opening_cash=cfg["opening_balances"]["1000"])
    write_bank_csv(rows, DATA_DIR / "sample_transactions_2026.csv")
    write_bank_xlsx(rows, DATA_DIR / "sample_transactions_2026.xlsx", cfg["business_name"])
    with open(DATA_DIR / "sample_truth_2026.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["txn_id", "category"])
        for r in rows:
            w.writerow([r["txn_id"], r["category"]])

    golden = build_golden()
    with open(DATA_DIR / "golden_transactions.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["txn_id", "date", "description", "amount", "category", "difficulty", "note"])
        w.writeheader()
        w.writerows(golden)

    scen_dir = DATA_DIR / "scenarios"
    scen_dir.mkdir(exist_ok=True)
    lean_cfg = json.loads(json.dumps(cfg))
    lean_cfg["business_name"] = "Kestrel Lean Studio LLC"
    lean_cfg["opening_balances"] = {"1000": 22000.00, "1010": 5000.00, "1500": 9000.00, "1510": 3000.00,
                                    "2100": 35000.00, "3000": 2000.00, "3900": -4000.00}
    lean_cfg["depreciation"]["opening_equipment_annual_depreciation"] = 1800.00
    for s in (
        scenario("northwind_full_year_2026", "Same business as the sample file, full calendar year", 2026, cfg, 1.0),
        scenario("lean_year_2026", "Smaller business, thinner margins, negative opening retained earnings",
                 4242, lean_cfg, 0.6),
    ):
        (scen_dir / f"{s['name']}.json").write_text(json.dumps(s, indent=1))

    cats = defaultdict(int)
    for r in rows:
        cats[r["category"]] += 1
    print(f"sample year: {len(rows)} transactions, ending balance {rows[-1]['balance']:,.2f}")
    print(f"golden set: {len(golden)} examples ({sum(g['difficulty'] == 'hard' for g in golden)} hard)")
    print("categories:", dict(sorted(cats.items())))


if __name__ == "__main__":
    main()
