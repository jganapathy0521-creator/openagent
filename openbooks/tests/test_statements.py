import json

import pandas as pd
import pytest

from books import DATA_DIR
from books.chart_of_accounts import ChartOfAccounts
from books.config import CompanyConfig
from books.ingest import ingest
from books.pipeline import run_pipeline
from evals.targets import OracleCategorizer

COA = ChartOfAccounts.load()
CFG = CompanyConfig.load()


@pytest.fixture(scope="module")
def result():
    df, report = ingest(DATA_DIR / "sample_transactions_2026.csv", CFG.period_start, CFG.period_end,
                        CFG.opening_cash)
    truth = dict(pd.read_csv(DATA_DIR / "sample_truth_2026.csv").values)
    return run_pipeline(df, CFG, COA, OracleCategorizer(truth), report=report)


def test_config_balances():
    assert CFG.validate(COA) == []


def test_all_checks_pass_with_correct_categories(result):
    for name, chk in result.statements["checks"].items():
        assert chk["passed"], f"{name}: {chk['detail']}"


def test_net_income_is_revenue_minus_expenses(result):
    pnl = result.statements["profit_and_loss"]
    totals = {s["name"]: s["total"] for s in pnl["sections"]}
    expected = (totals.get("Revenue", 0) + totals.get("Other Income", 0)
                - totals.get("Operating Expenses", 0) - totals.get("Other Expenses", 0))
    assert pnl["net_income"] == pytest.approx(expected, abs=0.01)


def test_balance_sheet_items_stay_off_the_pnl(result):
    pnl_accounts = {line["account"] for s in result.statements["profit_and_loss"]["sections"] for line in s["lines"]}
    for name in ("Owner Draws", "Owner Contributions", "Savings Account", "Loan Payable", "Equipment"):
        assert name not in pnl_accounts


def test_loan_split(result):
    sched = result.loan_schedule
    assert len(sched) == 12
    assert sched[0]["interest"] == pytest.approx(35000 * 0.075 / 12, abs=0.01)
    for row in sched:
        assert row["interest"] + row["principal"] == pytest.approx(row["payment"], abs=0.001)
    bs = {line["account"]: line["amount"] for s in result.statements["balance_sheet"]["sections"] for line in s["lines"]}
    assert bs["Loan Payable"] == pytest.approx(sched[-1]["balance_after"], abs=0.01)


def test_depreciation(result):
    pnl = {line["account"]: line["amount"] for s in result.statements["profit_and_loss"]["sections"] for line in s["lines"]}
    # 3,600 on opening equipment + laptop bought in May (7 of 36 months) + Dell bought in Sept (3 of 36)
    assert pnl["Depreciation Expense"] == pytest.approx(3600 + 2899 * 7 / 36 + 1640 * 3 / 36, abs=0.01)


def test_miscategorized_draw_changes_profit_not_balance(result):
    """Booking an owner draw as an expense still balances, but understates profit: exactly the
    kind of error the statement evals' line comparison is there to catch."""
    truth = dict(pd.read_csv(DATA_DIR / "sample_truth_2026.csv").values)
    draw_id = next(t for t, c in truth.items() if c == "Owner Draws")
    bad = {**truth, draw_id: "Uncategorized"}
    df = result.categorized[["txn_id", "date", "description", "amount", "balance"]].copy()
    r2 = run_pipeline(df, CFG, COA, OracleCategorizer(bad), report=result.report)
    assert r2.statements["checks"]["balance_sheet_balances"]["passed"]
    assert r2.statements["profit_and_loss"]["net_income"] < result.statements["profit_and_loss"]["net_income"]


def test_excel_export(result, tmp_path):
    from openpyxl import load_workbook

    from books.statements import export_excel
    path = tmp_path / "out.xlsx"
    export_excel(path, result.statements, result.categorized, result.journal_df, result.report.to_dict(),
                 result.loan_schedule)
    wb = load_workbook(path)
    assert {"Profit & Loss", "Balance Sheet", "Checks", "Transactions", "Review Queue", "Journal"} <= set(wb.sheetnames)
    json.dumps(result.statements)  # serializable
